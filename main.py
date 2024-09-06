import tkinter as tk
from docx import Document
from docx2pdf import convert
from datetime import datetime
import os

WORKING_DIR = os.getcwd()
WORD_DOC_LOC = os.path.join(WORKING_DIR, 'Zachary Kouba Cover Letter Template.docx')

def main():
    create_ui()

def create_ui():
    global entry_comp_name, entry_comp_address, entry_position, entry_aoi
    root = tk.Tk()
    root.title("Cover Letter Creation")

    tk.Label(root, text="Company Name").grid(row=0)
    entry_comp_name = tk.Entry(root)
    entry_comp_name.insert(0, "Test Name")
    entry_comp_name.grid(row=0, column=1)

    tk.Label(root, text="Company Address").grid(row=1)
    entry_comp_address = tk.Entry(root)
    entry_comp_address.insert(0, "Test Address")
    entry_comp_address.grid(row=1, column=1)

    tk.Label(root, text="Job Position").grid(row=2)
    entry_position = tk.Entry(root)
    entry_position.insert(0, "Test Position")
    entry_position.grid(row=2, column=1)

    tk.Label(root, text="Specific Area of Interest").grid(row=3)
    entry_aoi = tk.Entry(root)
    entry_aoi.insert(0, "Test AOI")
    entry_aoi.grid(row=3, column=1)

    submit_button = tk.Button(root, text="Create", command=submit_inputs)
    submit_button.grid(row=4, column=1)

    root.mainloop()

def submit_inputs():
    today_date = datetime.now().strftime("%B %d, %Y")
    comp_name = entry_comp_name.get()
    comp_address = entry_comp_address.get()
    position = entry_position.get()
    area_of_interest = entry_aoi.get()

    replacements = {
        "[Date]": today_date,
        "[Company Name]": comp_name,
        "[Company Location]": comp_address,
        "[Job Position]": position,
        "[AOI]": area_of_interest
    }
    input_to_word(replacements)

def input_to_word(replacements):
    doc = Document(WORD_DOC_LOC)
    replace_placeholders(doc, replacements)
    
    output_doc_name = f"Zachary Kouba Cover Letter {replacements['[Company Name]']}.docx"
    output_pdf_name = f"Zachary Kouba Cover Letter {replacements['[Company Name]']}.pdf"

    output_doc_filepath = os.path.join(WORKING_DIR, output_doc_name)
    output_pdf_filepath = os.path.join(WORKING_DIR, output_pdf_name)

    try:
        doc.save(output_doc_filepath)
        print(f"-> Docx Creation Successful at {output_doc_filepath}")
    except Exception as e:
        print(f"-> Docx Creation Error: {e}")

    try:
        convert(output_doc_filepath, output_pdf_filepath)
        print(f"-> PDF Creation Successful at {output_pdf_filepath}")
    except Exception as e:
        print(f"-> PDF Creation Error: {e}")

def replace_placeholders(doc, replacements):
    for paragraph in doc.paragraphs:
        # print(paragraph.text)
        for run in paragraph.runs:
            text = run.text
            for placeholder, value in replacements.items():
                text = text.replace(placeholder, value)
            run.text = text

if __name__ == "__main__":
    main()